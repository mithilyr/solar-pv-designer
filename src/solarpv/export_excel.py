"""Excel report: BOM, string schedule, inverter summary, checks and coloured layouts."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import __version__
from .layout import PANEL, SPECIAL
from .project import Project, ProjectResult

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CATEGORY_FILL = PatternFill("solid", start_color="D9E1F2")
MANUAL_FILL = PatternFill("solid", start_color="FFF2CC")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {
    "pass": PatternFill("solid", start_color="C6EFCE"),
    "warn": PatternFill("solid", start_color="FFEB9C"),
    "fail": PatternFill("solid", start_color="FFC7CE"),
    "info": PatternFill("solid", start_color="DDEBF7"),
}
STRING_COLOURS = ["F4B183", "A9D18E", "9DC3E6", "FFD966", "C9A0DC", "8FD3C8", "F8CBAD", "B4C7E7",
                  "E2EFDA", "FCE4D6", "D0CECE", "BDD7EE", "FFE699", "C5E0B4", "F4B084", "DBDBDB"]


def _header(ws, row: int, headers: list[str], widths: list[int] | None = None) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _title(ws, project: Project, title: str) -> int:
    ws["A1"] = f"{title} – {project.name}"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [f"Client: {project.client}" if project.client else "", f"Location: {project.location}" if project.location else "",
            f"Designer: {project.designer}" if project.designer else "",
            f"Generated {datetime.now():%Y-%m-%d %H:%M} by Solar PV Designer {__version__}"]
    ws["A2"] = "   ".join(m for m in meta if m)
    ws["A2"].font = Font(italic=True, size=9)
    return 4


def export_report(project: Project, result: ProjectResult, path: str | Path) -> None:
    wb = Workbook()
    _bom_sheet(wb.active, project, result)
    if result.design:
        _summary_sheet(wb.create_sheet("Inverters"), project, result)
        _schedule_sheet(wb.create_sheet("String Schedule"), project, result)
    if result.checks:
        _checks_sheet(wb.create_sheet("Design Checks"), project, result)
    for rr in result.roofs:
        _layout_sheet(wb.create_sheet(_sheet_name(f"Layout {rr.roof.name}", wb)), rr)
    wb.save(path)


def _sheet_name(name: str, wb) -> str:
    base = "".join(ch for ch in name if ch not in "[]:*?/\\")[:28]
    out, n = base, 2
    while out in wb.sheetnames:
        out, n = f"{base[:25]} {n}", n + 1
    return out


def _bom_sheet(ws, project: Project, result: ProjectResult) -> None:
    ws.title = "BOM"
    row = _title(ws, project, "Bill of Materials")
    if result.panel:
        ws.cell(row=3, column=1, value=f"{project.panels} panels × {result.panel.pmax_w:g} W = "
                                       f"{project.panels * result.panel.pmax_w / 1000:.2f} kWp").font = Font(size=10)
    _header(ws, row, ["S/N", "Description", "Code", "Unit", "Qty", "Notes"], [6, 58, 12, 7, 10, 70])
    row += 1
    category, n = None, 0
    for line in result.bom.lines:
        if line.category != category:
            category, n = line.category, 0
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell = ws.cell(row=row, column=1, value=category)
            cell.fill, cell.font = CATEGORY_FILL, Font(bold=True)
            row += 1
        n += 1
        values = [n, line.description, line.code, line.unit, line.qty, line.note]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = BORDER
            if line.qty is None:
                cell.fill = MANUAL_FILL
        row += 1
    row += 1
    ws.cell(row=row, column=2, value="Yellow rows are entered by hand.").font = Font(italic=True, size=9)
    ws.freeze_panes = "A5"


def _summary_sheet(ws, project: Project, result: ProjectResult) -> None:
    d, p = result.design, result.panel
    row = _title(ws, project, "Inverter Summary")
    ws.cell(row=3, column=1, value=f"DC {d.dc_kwp:.2f} kWp / AC {d.ac_kw:g} kW = {(d.dc_kwp / d.ac_kw - 1) * 100:.1f}% overload; "
                                   f"strings: " + ", ".join(f"{c} × {s} panels" for s, c in sorted(d.combination.items())))
    _header(ws, row, ["Inverter", "Model", "AC kW", "Strings", "Panels", "DC kWp", "Overload %", "MPPT layout"],
            [10, 32, 8, 9, 9, 10, 11, 70])
    for a in d.inverters:
        row += 1
        layout = "; ".join(f"MPPT{m.index}: {m.strings}×{m.string_size}" for m in a.mppts if m.strings)
        for c, v in enumerate([f"INV{a.number:02d}", a.inverter.key, a.inverter.ac_kw, a.strings, a.panels,
                               round(a.dc_kw(p), 2), round(a.overload_pct(p), 1), layout], 1):
            ws.cell(row=row, column=c, value=v).border = BORDER
    if d.warnings:
        row += 2
        for w in d.warnings:
            ws.cell(row=row, column=1, value=f"⚠ {w}").font = Font(color="9C5700")
            row += 1


def _schedule_sheet(ws, project: Project, result: ProjectResult) -> None:
    row = _title(ws, project, "String Schedule")
    _header(ws, row, ["String", "Inverter", "Model", "MPPT", "Panels in series", "Roof"], [16, 10, 32, 8, 16, 24])
    roof_of: dict[str, str] = {}
    for rr in result.roofs:
        for label in rr.string_map.values():
            roof_of.setdefault(label, rr.roof.name)
    for s in result.design.schedule():
        row += 1
        for c, v in enumerate([s["label"], f"INV{s['inverter']:02d}", s["model"], s["mppt"], s["panels"],
                               roof_of.get(s["label"], "")], 1):
            ws.cell(row=row, column=c, value=v).border = BORDER


def _checks_sheet(ws, project: Project, result: ProjectResult) -> None:
    row = _title(ws, project, "Design Checks")
    _header(ws, row, ["Check", "Result", "Detail", "Reference"], [42, 9, 100, 42])
    for ch in result.checks:
        row += 1
        for c, v in enumerate([ch.title, ch.status.upper(), ch.detail, ch.reference], 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=2).fill = STATUS_FILL[ch.status]


def _layout_sheet(ws, rr) -> None:
    grid = rr.roof.grid
    ws["A1"] = f"{rr.roof.name}: {rr.features['panels']} panels, roof type {rr.roof.roof_type}"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Suggested string routing follows the rail runs, column by column. Check against site conditions."
    ws["A2"].font = Font(italic=True, size=9)
    colour_of: dict[str, str] = {}
    for c in range(len(grid[0])):
        ws.column_dimensions[get_column_letter(c + 1)].width = 12
    for r, line in enumerate(grid):
        for c, v in enumerate(line):
            cell = ws.cell(row=r + 4, column=c + 1)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(size=8)
            if v == PANEL:
                label = rr.string_map.get((r, c), "")
                cell.value = label or "PV"
                colour = colour_of.setdefault(label, STRING_COLOURS[len(colour_of) % len(STRING_COLOURS)]) if label else "C6EFCE"
                cell.fill = PatternFill("solid", start_color=colour)
            elif v == SPECIAL:
                cell.value = "2"
                cell.fill = PatternFill("solid", start_color="F8CBAD")
