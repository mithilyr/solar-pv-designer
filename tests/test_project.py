from openpyxl import load_workbook

from solarpv.export_excel import export_report
from solarpv.project import InverterChoice, Project, Roof, compute

GRID = [[1] * 12 for _ in range(12)] + [[0] * 12] + [[1] * 12 for _ in range(12)]


def make_project():
    return Project(name="Test site", panel_key="JA Solar JAM72D40-595/LB",
                   inverters=[InverterChoice("SolaX X3-FTH-75K", 2)],
                   roofs=[Roof("Main", [row[:] for row in GRID]), Roof("Annex", [[1] * 6 for _ in range(6)], "clay_tile")])


def test_pipeline_and_excel(tmp_path, library, rules):
    project = make_project()
    result = compute(project, library, rules)
    assert result.design is not None, result.design_error
    assert result.design.panels_used == project.panels == 324
    assert sum(len(r.string_map) for r in result.roofs) == 324
    names = [l.description for l in result.bom.lines]
    assert any(n.startswith("INVERTER SolaX X3-FTH-75K") for n in names)
    assert any("TILE ROOF HOOK" in n for n in names)
    assert result.site_quantities["dc_cable_4mm_m"] + result.site_quantities["dc_cable_6mm_m"] > 0
    assert any(c.title == "String fuses" for c in result.checks)

    path = tmp_path / "report.xlsx"
    export_report(project, result, path)
    wb = load_workbook(path)
    assert wb.sheetnames[:4] == ["BOM", "Inverters", "String Schedule", "Design Checks"]


def test_manual_quantities_and_save(tmp_path, library, rules):
    project = make_project()
    project.manual_quantities["EARTH ROD 10FT"] = 4
    result = compute(project, library, rules)
    rod = next(l for l in result.bom.lines if l.description == "EARTH ROD 10FT")
    assert rod.qty == 4
    path = tmp_path / "site.spvd"
    project.save(path)
    again = Project.load(path)
    assert again.panels == project.panels
    assert again.manual_quantities == {"EARTH ROD 10FT": 4}
    assert again.inverters[0].key == "SolaX X3-FTH-75K"


def test_library_integrity(library):
    for inv in library.inverters:
        assert not inv.validate(), (inv.key, inv.validate())
    assert len({i.key for i in library.inverters}) == len(library.inverters)


def test_unconnected_panels_are_a_check_warning(library, rules):
    project = Project(name="House", panel_key="JA Solar JAM72D40-595/LB",
                      inverters=[InverterChoice("SolaX X1-BOOST-5K-G4", 1)],
                      roofs=[Roof("Roof", [[1] * 7, [1] * 7])])
    result = compute(project, library, rules)
    warn = [c for c in result.checks if c.title == "Unconnected panels"]
    assert warn and warn[0].status == "warn" and warn[0].detail.startswith("3 of 14")


def test_excel_bom_leaves_out_zero_lines(tmp_path, library, rules):
    project = make_project()
    result = compute(project, library, rules)
    assert any(line.qty == 0 for line in result.bom.lines)
    export_report(project, result, tmp_path / "r.xlsx")
    ws = load_workbook(tmp_path / "r.xlsx")["BOM"]
    assert all(row[4] != 0 for row in ws.iter_rows(min_row=5, values_only=True) if row[1])
